"""
PR-RAS-NPF i BIL-NPF XLSX generator za FINA.
Generira izvještaje prema Pravilniku o izvještavanju u neprofitnom računovodstvu.

Usage:
    from udruge.fina_reports import generate_prras_npf, generate_bil_npf
    xlsx_bytes = generate_prras_npf(company, year=2026, period='yearly')
"""
import io
from datetime import date
from decimal import Decimal

from django.db.models import Sum, Q
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill


# ═══════════════════════════════════════════════════════════════════════════
# PR-RAS-NPF: AOP mapping (konto -> AOP)
# Stupac 4 = prethodna godina, Stupac 5 = tekuce razdoblje
# ═══════════════════════════════════════════════════════════════════════════

PRRAS_ROWS = [
    # (AOP, konto_prefix, opis, formula_or_None)
    # PRIHODI
    (1,  '3',     'PRIHODI (AOP 002+005+008+011+024+032+041)', 'sum:2,5,8,11,24,32,41'),
    (2,  '31',    'Prihodi od prodaje roba i pruzanja usluga (AOP 003+004)', 'sum:3,4'),
    (3,  '3111',  'Prihodi od prodaje roba', None),
    (4,  '3112',  'Prihodi od pruzanja usluga', None),
    (5,  '32',    'Prihodi od clanarina i clanskih doprinosa (AOP 006+007)', 'sum:6,7'),
    (6,  '3211',  'Clanarine', None),
    (7,  '3212',  'Clanski doprinosi', None),
    (8,  '33',    'Prihodi po posebnim propisima (AOP 009+010)', 'sum:9,10'),
    (9,  '3311',  'Prihodi po posebnim propisima iz proracuna', None),
    (10, '3312',  'Prihodi po posebnim propisima iz ostalih izvora', None),
    (11, '34',    'Prihodi od imovine (AOP 012+021)', 'sum:12,21'),
    (12, '341',   'Prihodi od financijske imovine (AOP 013 do 020)', 'sum:13,14,15,16,17,18,19,20'),
    (13, '3411',  'Prihodi od kamata za dane zajmove', None),
    (14, '3412',  'Prihodi od kamata po vrijednosnim papirima', None),
    (15, '3413',  'Kamate na orocena sredstva i depozite po vidjenju', None),
    (16, '3414',  'Prihodi od zateznih kamata', None),
    (17, '3415',  'Prihodi od pozitivnih tecajnih razlika', None),
    (18, '3416',  'Prihodi od dividendi', None),
    (19, '3417',  'Prihodi od dobiti trgovackih drustava', None),
    (20, '3418',  'Ostali prihodi od financijske imovine', None),
    (21, '342',   'Prihodi od nefinancijske imovine (AOP 022+023)', 'sum:22,23'),
    (22, '3421',  'Prihodi od zakupa i iznajmljivanja imovine', None),
    (23, '3422',  'Ostali prihodi od nefinancijske imovine', None),
    (24, '35',    'Prihodi od donacija (AOP 025+028+029+030+031)', 'sum:25,28,29,30,31'),
    (25, '351',   'Prihodi od donacija iz proracuna (AOP 026+027)', 'sum:26,27'),
    (26, '3511',  'Prihodi od donacija iz drzavnog proracuna', None),
    (27, '3512',  'Prihodi od donacija iz proracuna JLP(R)S', None),
    (28, '352',   'Prihodi od inozemnih vlada i medunarodnih organizacija', None),
    (29, '353',   'Prihodi od trgovackih drustava i ostalih pravnih osoba', None),
    (30, '354',   'Prihodi od gradana i kucanstava', None),
    (31, '355',   'Ostali prihodi od donacija', None),
    (32, '36',    'Ostali prihodi (AOP 033+036+037)', 'sum:33,36,37'),
    (33, '361',   'Prihodi od naknade stete i refundacija (AOP 034+035)', 'sum:34,35'),
    (34, '3611',  'Prihodi od naknade steta', None),
    (35, '3612',  'Prihodi od refundacija', None),
    (36, '362',   'Prihodi od prodaje dugotrajne imovine', None),
    (37, '363',   'Ostali nespomenuti prihodi (AOP 038+039+040)', 'sum:38,39,40'),
    (38, '3631',  'Otpis obveza', None),
    (39, '3632',  'Naplacena otpisana potrazivanja', None),
    (40, '3633',  'Ostali nespomenuti prihodi', None),
    (41, '37',    'Prihodi od povezanih neprofitnih organizacija (AOP 042+043)', 'sum:42,43'),
    (42, '3711',  'Tekuci prihodi od povezanih neprofitnih organizacija', None),
    (43, '3712',  'Kapitalni prihodi od povezanih neprofitnih organizacija', None),
    # RASHODI
    (44, '4',     'RASHODI (AOP 045+057+098+099+110+115+126)', 'sum:45,57,98,99,110,115,126'),
    (45, '41',    'Rashodi za radnike (AOP 046+051+052)', 'sum:46,51,52'),
    (46, '411',   'Place (AOP 047+048+049+050)', 'sum:47,48,49,50'),
    (47, '4111',  'Place za redovan rad', None),
    (48, '4112',  'Place u naravi', None),
    (49, '4113',  'Place za prekovremeni rad', None),
    (50, '4114',  'Place za posebne uvjete rada', None),
    (51, '412',   'Ostali rashodi za radnike', None),
    (52, '413',   'Doprinosi na place (AOP 053+054+055+056)', 'sum:53,54,55,56'),
    (53, '4131',  'Doprinosi za zdravstveno osiguranje', None),
    (54, '4132',  'Doprinosi za zaposljavanje', None),
    (55, '4133',  'Doprinosi za mirovinsko koje placa poslodavac', None),
    (56, '4134',  'Posebni doprinos za poticanje zaposljavanja OSI', None),
    (57, '42',    'Materijalni rashodi (AOP 058+062+067+072+077+087+092)', 'sum:58,62,67,72,77,87,92'),
    (58, '421',   'Naknade troskova radnicima (AOP 059+060+061)', 'sum:59,60,61'),
    (59, '4211',  'Sluzbena putovanja', None),
    (60, '4212',  'Naknade za prijevoz, rad na terenu i odvojeni zivot', None),
    (61, '4213',  'Strucno usavrsavanje radnika', None),
    (62, '422',   'Naknade clanovima u predstavnickim tijelima (AOP 063-066)', 'sum:63,64,65,66'),
    (63, '4221',  'Naknade za obavljanje aktivnosti', None),
    (64, '4222',  'Naknade troskova sluzbenih putovanja', None),
    (65, '4223',  'Naknade ostalih troskova', None),
    (66, '4224',  'Ostale naknade', None),
    (67, '423',   'Naknade volonterima (AOP 068+069+070+071)', 'sum:68,69,70,71'),
    (68, '4231',  'Naknade za obavljanje djelatnosti', None),
    (69, '4232',  'Naknade troskova sluzbenih putovanja', None),
    (70, '4233',  'Naknade ostalih troskova', None),
    (71, '4234',  'Ostale naknade', None),
    (72, '424',   'Naknade ostalim osobama izvan radnog odnosa (AOP 073-076)', 'sum:73,74,75,76'),
    (73, '4241',  'Naknade za obavljanje aktivnosti', None),
    (74, '4242',  'Naknade troskova sluzbenih putovanja', None),
    (75, '4243',  'Naknade ostalih troskova', None),
    (76, '4244',  'Ostale naknade', None),
    (77, '425',   'Rashodi za usluge (AOP 078-086)', 'sum:78,79,80,81,82,83,84,85,86'),
    (78, '4251',  'Usluge telefona, poste i prijevoza', None),
    (79, '4252',  'Usluge tekuceg i investicijskog odrzavanja', None),
    (80, '4253',  'Usluge promidzbe i informiranja', None),
    (81, '4254',  'Komunalne usluge', None),
    (82, '4255',  'Zakupnine i najamnine', None),
    (83, '4256',  'Zdravstvene i veterinarske usluge', None),
    (84, '4257',  'Intelektualne i osobne usluge', None),
    (85, '4258',  'Racunalne usluge', None),
    (86, '4259',  'Ostale usluge', None),
    (87, '426',   'Rashodi za materijal i energiju (AOP 088-091)', 'sum:88,89,90,91'),
    (88, '4261',  'Uredski materijal i ostali materijalni rashodi', None),
    (89, '4262',  'Materijal i sirovine', None),
    (90, '4263',  'Energija', None),
    (91, '4264',  'Sitan inventar i auto gume', None),
    (92, '429',   'Ostali nespomenuti materijalni rashodi (AOP 093-097)', 'sum:93,94,95,96,97'),
    (93, '4291',  'Premije osiguranja', None),
    (94, '4292',  'Reprezentacija', None),
    (95, '4293',  'Clanarine (rashod)', None),
    (96, '4294',  'Kotizacije', None),
    (97, '4295',  'Ostali nespomenuti materijalni rashodi', None),
    (98, '43',    'Rashodi amortizacije', None),
    (99, '44',    'Financijski rashodi (AOP 100+101+105)', 'sum:100,101,105'),
    (100, '441',  'Kamate za izdane vrijednosne papire', None),
    (101, '442',  'Kamate za primljene kredite i zajmove (AOP 102-104)', 'sum:102,103,104'),
    (102, '4421', 'Kamate za primljene kredite banaka', None),
    (103, '4422', 'Kamate za primljene robne i ostale zajmove', None),
    (104, '4423', 'Kamate za odobrene a nerealizirane kredite', None),
    (105, '443',  'Ostali financijski rashodi (AOP 106-109)', 'sum:106,107,108,109'),
    (106, '4431', 'Bankarske usluge i usluge platnog prometa', None),
    (107, '4432', 'Negativne tecajne razlike i valutna klauzula', None),
    (108, '4433', 'Zatezne kamate', None),
    (109, '4434', 'Ostali nespomenuti financijski rashodi', None),
    (110, '45',   'Donacije (AOP 111+114)', 'sum:111,114'),
    (111, '451',  'Tekuce donacije (AOP 112+113)', 'sum:112,113'),
    (112, '4511', 'Tekuce donacije', None),
    (113, '4512', 'Stipendije', None),
    (114, '452',  'Kapitalne donacije', None),
    (115, '46',   'Ostali rashodi (AOP 116+121)', 'sum:116,121'),
    (116, '461',  'Kazne, penali i naknade stete (AOP 117-120)', 'sum:117,118,119,120'),
    (117, '4611', 'Naknade steta pravnim i fizickim osobama', None),
    (118, '4612', 'Penali, lezarine i drugo', None),
    (119, '4613', 'Naknade steta radnicima', None),
    (120, '4614', 'Ugovorene kazne i ostale naknade steta', None),
    (121, '462',  'Ostali nespomenuti rashodi (AOP 122-125)', 'sum:122,123,124,125'),
    (122, '4621', 'Neotpisana vrijednost i drugi rashodi otudene imovine', None),
    (123, '4622', 'Otpisana potrazivanja', None),
    (124, '4623', 'Rashodi za ostala porezna davanja', None),
    (125, '4624', 'Ostali nespomenuti rashodi', None),
    (126, '47',   'Rashodi vezani uz financ. povezanih NPF org (AOP 127+128)', 'sum:127,128'),
    (127, '4711', 'Tekuci rashodi vezani uz financ. povezanih NPF org', None),
    (128, '4712', 'Kapitalni rashodi vezani uz financ. povezanih NPF org', None),
    # Zalihe
    (129, None,   'Stanje zaliha na pocetku razdoblja', 'special'),
    (130, None,   'Stanje zaliha na kraju razdoblja', 'special'),
    (131, None,   'Povecanje zaliha (AOP 130-129)', 'calc:130-129'),
    (132, None,   'Smanjenje zaliha (AOP 129-130)', 'calc:129-130'),
    (133, None,   'UKUPNI RASHODI (AOP 044-131 ili 044+132)', 'special'),
    (134, None,   'VISAK PRIHODA (AOP 001-133)', 'calc:1-133'),
    (135, None,   'MANJAK PRIHODA (AOP 133-001)', 'calc:133-1'),
    (136, '5221', 'Visak prihoda - preneseni', None),
    (137, '5222', 'Manjak prihoda - preneseni', None),
    (138, None,   'Visak prihoda raspoloziv u sljedecem razdoblju', 'calc:134-135+136-137'),
    (139, None,   'Manjak prihoda za pokrice u sljedecem razdoblju', 'calc:135+137-134-136'),
    # Dodatni podaci
    (140, '11',   'Stanje novcanih sredstava na pocetku godine', None),
    (141, None,   'Ukupni priljevi na novcane racune i blagajne', 'special'),
    (142, None,   'Ukupni odljevi s novcanih racuna i blagajni', 'special'),
    (143, None,   'Stanje novcanih sredstava na kraju razdoblja', 'calc:140+141-142'),
    (144, None,   'Prosjecan broj radnika na osnovi stanja', 'special'),
    (145, None,   'Prosjecan broj radnika na osnovi sati rada', 'special'),
]


def _get_konto_balance(company, konto_prefix, year, period_end_month=12):
    """Dohvati saldo za konto prefix iz JournalEntry/LedgerLine."""
    from konto.models import KontoAccount
    from django.db.models import Sum

    # Dohvati sve konta koja pocnu s prefixom
    accounts = KontoAccount.objects.filter(
        company=company, code__startswith=konto_prefix, is_active=True,
    )
    if not accounts.exists():
        return Decimal('0')

    # Dohvati ledger linije za period
    try:
        from konto.models import LedgerLine
        lines = LedgerLine.objects.filter(
            account__in=accounts,
            journal_entry__date__year=year,
            journal_entry__date__month__lte=period_end_month,
        )
        debits = lines.filter(side='debit').aggregate(t=Sum('amount'))['t'] or Decimal('0')
        credits = lines.filter(side='credit').aggregate(t=Sum('amount'))['t'] or Decimal('0')

        # Prihodi (klasa 3) rastu na potraznoj, rashodi (klasa 4) na dugovnoj
        if konto_prefix.startswith('3') or konto_prefix.startswith('5'):
            return credits - debits
        else:
            return debits - credits
    except (ValueError, TypeError):
        return Decimal('0')


def generate_prras_npf(company, year=None, period='yearly'):
    """
    Generira PR-RAS-NPF izvjestaj u XLSX formatu.

    Args:
        company: Company objekt (mora biti nonprofit)
        year: godina izvjestaja (default: tekuca)
        period: 'yearly' (cijela godina) ili 'half' (polugodisnji)

    Returns:
        bytes: XLSX sadrzaj
    """
    if year is None:
        year = date.today().year

    period_end_month = 12 if period == 'yearly' else 6
    prev_year = year - 1

    wb = Workbook()
    ws = wb.active
    ws.title = 'PR-RAS-NPF'

    # Stilovi
    bold = Font(bold=True, size=10)
    header_font = Font(bold=True, size=12)
    section_fill = PatternFill(start_color='F2F2F2', end_color='F2F2F2', fill_type='solid')
    border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin'),
    )

    # Zaglavlje
    ws.merge_cells('A1:J1')
    ws['A1'] = 'IZVJESTAJ O PRIHODIMA I RASHODIMA NEPROFITNIH ORGANIZACIJA'
    ws['A1'].font = header_font
    ws['A1'].alignment = Alignment(horizontal='center')

    period_label = '1. sijecnja do 31. prosinca' if period == 'yearly' else '1. sijecnja do 30. lipnja'
    ws['A2'] = 'za razdoblje %s %d.' % (period_label, year)
    ws['A3'] = 'Naziv obveznika: %s' % company.name
    ws['A4'] = 'OIB: %s' % company.oib
    ws['A5'] = 'Adresa: %s, %s %s' % (company.address, company.zip_code, company.city)
    ws['A6'] = 'Obrazac PR-RAS-NPF'
    ws['H6'] = 'AOP oznaka razdoblja: %d-%02d' % (year, period_end_month)

    # Zaglavlje tablice
    row = 8
    headers = ['Racun iz rac. plana', 'OPIS', '', '', '', '', '', 'AOP',
               'Ostvareno prethodne godine', 'Ostvareno tekuce razdoblje', 'Index (5/4)']
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.font = bold
        cell.border = border
        cell.fill = section_fill

    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=7)

    # Podaci
    aop_values = {}  # AOP -> (prev, curr) za formule

    # Prvo izracunaj sve leaf AOP-ove (bez formula)
    for aop, konto, opis, formula in PRRAS_ROWS:
        if konto and not formula:
            prev = _get_konto_balance(company, konto, prev_year, 12 if period == 'yearly' else 6)
            curr = _get_konto_balance(company, konto, year, period_end_month)
            aop_values[aop] = (prev, curr)

    # Izracunaj formula AOP-ove
    for aop, konto, opis, formula in PRRAS_ROWS:
        if formula and formula.startswith('sum:'):
            refs = [int(x) for x in formula[4:].split(',')]
            prev = sum(aop_values.get(r, (Decimal('0'), Decimal('0')))[0] for r in refs)
            curr = sum(aop_values.get(r, (Decimal('0'), Decimal('0')))[1] for r in refs)
            aop_values[aop] = (prev, curr)
        elif formula and formula.startswith('calc:'):
            # Simple calc: e.g. "1-133" or "134-135+136-137"
            pass  # Will calculate after all sums
        elif formula == 'special':
            aop_values[aop] = (Decimal('0'), Decimal('0'))

    # Calc formulas (second pass)
    for aop, konto, opis, formula in PRRAS_ROWS:
        if formula and formula.startswith('calc:'):
            expr = formula[5:]
            # Parse simple +/- expression
            prev_val = Decimal('0')
            curr_val = Decimal('0')
            parts = expr.replace('-', '+-').split('+')
            for p in parts:
                p = p.strip()
                if not p:
                    continue
                neg = p.startswith('-')
                ref = int(p.lstrip('-'))
                pv, cv = aop_values.get(ref, (Decimal('0'), Decimal('0')))
                if neg:
                    prev_val -= pv
                    curr_val -= cv
                else:
                    prev_val += pv
                    curr_val += cv
            # Visak/manjak: samo pozitivne vrijednosti
            if aop in (134, 138):
                prev_val = max(Decimal('0'), prev_val)
                curr_val = max(Decimal('0'), curr_val)
            elif aop in (135, 139):
                prev_val = max(Decimal('0'), prev_val)
                curr_val = max(Decimal('0'), curr_val)
            aop_values[aop] = (prev_val, curr_val)

    # Ispis redaka
    for aop, konto, opis, formula in PRRAS_ROWS:
        row += 1
        prev, curr = aop_values.get(aop, (Decimal('0'), Decimal('0')))

        # Sekcija header
        is_section = formula and formula.startswith('sum:')
        is_total = aop in (1, 44, 133, 134, 135, 138, 139)

        ws.cell(row=row, column=1, value=konto or '').border = border
        cell_opis = ws.cell(row=row, column=2, value=opis)
        cell_opis.border = border
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=7)

        ws.cell(row=row, column=8, value=aop).border = border
        ws.cell(row=row, column=9, value=float(prev)).border = border
        ws.cell(row=row, column=9).number_format = '#,##0'
        ws.cell(row=row, column=10, value=float(curr)).border = border
        ws.cell(row=row, column=10).number_format = '#,##0'

        # Index
        if prev and prev != 0:
            idx = int(curr / prev * 100)
            ws.cell(row=row, column=11, value=idx).border = border
        else:
            ws.cell(row=row, column=11, value='-').border = border

        if is_total or is_section:
            for c in range(1, 12):
                ws.cell(row=row, column=c).font = bold
                if is_total:
                    ws.cell(row=row, column=c).fill = section_fill

    # Sirine stupaca
    ws.column_dimensions['A'].width = 12
    ws.column_dimensions['B'].width = 55
    ws.column_dimensions['H'].width = 6
    ws.column_dimensions['I'].width = 18
    ws.column_dimensions['J'].width = 18
    ws.column_dimensions['K'].width = 10

    # Save
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


# ═══════════════════════════════════════════════════════════════════════════
# G-PR-IZ-NPF: Godisnji izvjestaj o primicima i izdacima
# Za udruge s jednostavnim racunovodstvom (< 30.000 EUR prihoda)
# ═══════════════════════════════════════════════════════════════════════════

GPRIZNPF_ROWS = [
    (1, 'PRIMICI', 'primitak', None, 'sum:2,3,4,5,6,7'),
    (2, 'Clanarine i clanski doprinosi', 'primitak', 'clanarina', None),
    (3, 'Donacije i potpore iz proracuna', 'primitak', 'dotacija', None),
    (4, 'Donacije od pravnih i fizickih osoba', 'primitak', 'donacija', None),
    (5, 'Prihodi od gospodarske djelatnosti', 'primitak', 'prodaja', None),
    (6, 'Prihodi od imovine', 'primitak', None, None),
    (7, 'Ostali primici', 'primitak', 'ostalo', None),
    (8, 'IZDACI', 'izdatak', None, 'sum:9,10,11,12,13,14'),
    (9, 'Place i naknade', 'izdatak', 'placa', None),
    (10, 'Materijalni rashodi', 'izdatak', 'materijal', None),
    (11, 'Rezije (komunalne usluge, energija)', 'izdatak', 'rezije', None),
    (12, 'Usluge (vanjski suradnici, prijevoz)', 'izdatak', None, None),
    (13, 'Donacije dane', 'izdatak', None, None),
    (14, 'Ostali izdaci', 'izdatak', 'ostalo', None),
    (15, 'VISAK PRIMITAKA (r1-r8)', None, None, 'calc:1-8'),
    (16, 'MANJAK PRIMITAKA (r8-r1)', None, None, 'calc:8-1'),
    (17, 'Stanje novca na pocetku godine', None, None, 'special'),
    (18, 'Stanje novca na kraju godine (r17+r15-r16)', None, None, 'calc:17+15-16'),
]


def generate_gpriznpf(company, year=None):
    """
    Generira G-PR-IZ-NPF izvjestaj za jednostavno racunovodstvo.
    Povlaci podatke iz KnjigaPrimitaka modela.
    """
    if year is None:
        year = date.today().year

    from udruge.models import KnjigaPrimitaka

    wb = Workbook()
    ws = wb.active
    ws.title = 'G-PR-IZ-NPF'

    bold = Font(bold=True, size=10)
    header_font = Font(bold=True, size=12)
    section_fill = PatternFill(start_color='F2F2F2', end_color='F2F2F2', fill_type='solid')
    border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin'),
    )

    ws.merge_cells('A1:F1')
    ws['A1'] = 'GODISNJI FINANCIJSKI IZVJESTAJ O PRIMICIMA I IZDACIMA'
    ws['A1'].font = header_font
    ws['A2'] = 'za razdoblje 1. sijecnja do 31. prosinca %d.' % year
    ws['A3'] = 'Naziv: %s' % company.name
    ws['A4'] = 'OIB: %s' % company.oib
    ws['A5'] = 'Obrazac G-PR-IZ-NPF'

    row = 7
    for col, h in enumerate(['R.br.', 'OPIS', 'Prethodna godina', 'Tekuca godina', 'Index'], 1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.font = bold
        cell.border = border
        cell.fill = section_fill

    # Dohvati podatke iz KnjigaPrimitaka
    qs = KnjigaPrimitaka.objects.filter(company=company, datum__year=year)
    qs_prev = KnjigaPrimitaka.objects.filter(company=company, datum__year=year - 1)

    aop_values = {}
    for rb, opis, vrsta, izvor, formula in GPRIZNPF_ROWS:
        if vrsta and not formula:
            f = {'vrsta': vrsta}
            if izvor:
                f['izvor'] = izvor
            curr = qs.filter(**f).aggregate(s=Sum('iznos'))['s'] or Decimal('0')
            prev = qs_prev.filter(**f).aggregate(s=Sum('iznos'))['s'] or Decimal('0')
            aop_values[rb] = (prev, curr)

    # Formule
    for rb, opis, vrsta, izvor, formula in GPRIZNPF_ROWS:
        if formula and formula.startswith('sum:'):
            refs = [int(x) for x in formula[4:].split(',')]
            prev = sum(aop_values.get(r, (Decimal('0'), Decimal('0')))[0] for r in refs)
            curr = sum(aop_values.get(r, (Decimal('0'), Decimal('0')))[1] for r in refs)
            aop_values[rb] = (prev, curr)
        elif formula and formula.startswith('calc:'):
            expr = formula[5:]
            prev_val = Decimal('0')
            curr_val = Decimal('0')
            parts = expr.replace('-', '+-').split('+')
            for p in parts:
                p = p.strip()
                if not p:
                    continue
                neg = p.startswith('-')
                ref = int(p.lstrip('-'))
                pv, cv = aop_values.get(ref, (Decimal('0'), Decimal('0')))
                if neg:
                    prev_val -= pv
                    curr_val -= cv
                else:
                    prev_val += pv
                    curr_val += cv
            if rb in (15,):
                prev_val = max(Decimal('0'), prev_val)
                curr_val = max(Decimal('0'), curr_val)
            elif rb in (16,):
                prev_val = max(Decimal('0'), prev_val)
                curr_val = max(Decimal('0'), curr_val)
            aop_values[rb] = (prev_val, curr_val)
        elif formula == 'special':
            aop_values[rb] = (Decimal('0'), Decimal('0'))

    for rb, opis, vrsta, izvor, formula in GPRIZNPF_ROWS:
        row += 1
        prev, curr = aop_values.get(rb, (Decimal('0'), Decimal('0')))
        is_total = rb in (1, 8, 15, 16, 18)

        ws.cell(row=row, column=1, value=rb).border = border
        ws.cell(row=row, column=2, value=opis).border = border
        ws.cell(row=row, column=3, value=float(prev)).border = border
        ws.cell(row=row, column=3).number_format = '#,##0'
        ws.cell(row=row, column=4, value=float(curr)).border = border
        ws.cell(row=row, column=4).number_format = '#,##0'
        idx = int(curr / prev * 100) if prev and prev != 0 else 0
        ws.cell(row=row, column=5, value=idx if idx else '-').border = border

        if is_total:
            for c in range(1, 6):
                ws.cell(row=row, column=c).font = bold
                ws.cell(row=row, column=c).fill = section_fill

    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 45
    ws.column_dimensions['C'].width = 18
    ws.column_dimensions['D'].width = 18
    ws.column_dimensions['E'].width = 10

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()



# ═══════════════════════════════════════════════════════════════════════════
# BIL-NPF: Bilanca neprofitnih organizacija
# Stanje imovine, obveza i vlastitih izvora na 1.1. i 31.12.
# ═══════════════════════════════════════════════════════════════════════════

BIL_ROWS = [
    # (AOP, konto_prefix, opis, formula_or_None)
    # AKTIVA
    (1,   None,  "IMOVINA (AOP 002+074)", "sum:2,74"),
    (2,   None,  "Nefinancijska imovina (AOP 003+018+047+051+055+064)", "sum:3,18,47,51,55,64"),
    (3,   "01",  "Neproizvedena dugotrajna imovina (AOP 004 do 017)", None),
    (18,  "02",  "Proizvedena dugotrajna imovina (AOP 019 do 046)", None),
    (47,  "03",  "Plemeniti metali i ostale pohranjene vrijednosti (AOP 048 do 050)", None),
    (51,  "04",  "Sitni inventar (AOP 052 do 054)", None),
    (55,  "05",  "Nefinancijska imovina u pripremi (AOP 056 do 063)", None),
    (64,  "06",  "Proizvedena kratkotrajna imovina (AOP 065 do 073)", None),
    (74,  None,  "Financijska imovina (AOP 075+083+100+105)", "sum:75,83,100,105"),
    (75,  "11",  "Novac u banci i blagajni (AOP 076 do 082)", None),
    (83,  None,  "Depoziti, jamcevni polozi i potrazivanja (AOP 084 do 099)", "sum_konto:12,13,16"),
    (100, "14",  "Zajmovi (AOP 101 do 104)", None),
    (105, "15",  "Vrijednosni papiri (AOP 106 do 109)", None),
    # PASIVA
    (146, None,  "OBVEZE I VLASTITI IZVORI (AOP 147+163)", "sum:147,163"),
    (147, None,  "Obveze (AOP 148+153+158+162)", "sum:148,153,158,162"),
    (148, None,  "Obveze za rashode (AOP 149 do 152)", "sum_konto:21,22"),
    (153, "23",  "Obveze za kredite i zajmove (AOP 154 do 157)", None),
    (158, "24",  "Odgodeno placanje rashoda i prihodi buducih razdoblja (AOP 159 do 161)", None),
    (162, "25",  "Obveze za poreze i ostale obveze (AOP 162a do 162d)", None),
    (163, None,  "Vlastiti izvori (AOP 164+168)", "sum:164,168"),
    (164, "51",  "Vlastiti izvori (AOP 165 do 167)", None),
    (168, "52",  "Visak/manjak prihoda (AOP 169 do 175)", None),
]


def _get_konto_balance_at_date(company, konto_prefix, year, month, day):
    """
    Dohvati kumulativni saldo za konto prefix do zadanog datuma.
    Za bilancu trebamo stanje na 1.1. (sve do 31.12. prethodne) i na 31.12. tekuce.
    """
    from konto.models import KontoAccount
    from django.db.models import Sum

    accounts = KontoAccount.objects.filter(
        company=company, code__startswith=konto_prefix, is_active=True,
    )
    if not accounts.exists():
        return Decimal("0")

    try:
        from konto.models import LedgerLine
        target_date = date(year, month, day)
        lines = LedgerLine.objects.filter(
            account__in=accounts,
            journal_entry__date__lte=target_date,
        )
        debits = lines.filter(side="debit").aggregate(t=Sum("amount"))["t"] or Decimal("0")
        credits = lines.filter(side="credit").aggregate(t=Sum("amount"))["t"] or Decimal("0")

        # Aktiva (klase 0,1) raste na dugovnoj, pasiva (klase 2,5) na potraznoj
        if konto_prefix.startswith(("2", "5")):
            return credits - debits
        else:
            return debits - credits
    except (ValueError, TypeError):
        return Decimal("0")


def generate_bil_npf(company, year=None):
    """
    Generira BIL-NPF (Bilanca neprofitnih organizacija) u XLSX formatu.

    Args:
        company: Company objekt (mora biti nonprofit)
        year: godina izvjestaja (default: tekuca)

    Returns:
        bytes: XLSX sadrzaj
    """
    if year is None:
        year = date.today().year

    wb = Workbook()
    ws = wb.active
    ws.title = "BIL-NPF"

    # Stilovi
    bold = Font(bold=True, size=10)
    header_font = Font(bold=True, size=12)
    section_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    # Zaglavlje
    ws.merge_cells("A1:H1")
    ws["A1"] = "BILANCA - NEPROFITNE ORGANIZACIJE"
    ws["A1"].font = header_font
    ws["A1"].alignment = Alignment(horizontal="center")

    ws["A2"] = "na dan 31. prosinca %d." % year
    ws["A3"] = "Naziv obveznika: %s" % company.name
    ws["A4"] = "OIB: %s" % company.oib
    ws["A5"] = "Adresa: %s, %s %s" % (company.address, company.zip_code, company.city)
    ws["A6"] = "Obrazac BIL-NPF"

    # Zaglavlje tablice
    row = 8
    headers = ["Racun iz rac. plana", "OPIS", "", "", "", "AOP",
               "Stanje 1. sijecnja", "Stanje 31. prosinca", "Index"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.font = bold
        cell.border = border
        cell.fill = section_fill

    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)

    # Izracunaj leaf AOP-ove (bez formula)
    aop_values = {}

    for aop, konto, opis, formula in BIL_ROWS:
        if konto and not formula:
            opening = _get_konto_balance_at_date(company, konto, year - 1, 12, 31)
            closing = _get_konto_balance_at_date(company, konto, year, 12, 31)
            aop_values[aop] = (opening, closing)
        elif formula and formula.startswith("sum_konto:"):
            prefixes = formula[10:].split(",")
            opening = Decimal("0")
            closing = Decimal("0")
            for pfx in prefixes:
                opening += _get_konto_balance_at_date(company, pfx.strip(), year - 1, 12, 31)
                closing += _get_konto_balance_at_date(company, pfx.strip(), year, 12, 31)
            aop_values[aop] = (opening, closing)

    # Izracunaj formula AOP-ove
    for aop, konto, opis, formula in BIL_ROWS:
        if formula and formula.startswith("sum:"):
            refs = [int(x) for x in formula[4:].split(",")]
            opening = sum(aop_values.get(r, (Decimal("0"), Decimal("0")))[0] for r in refs)
            closing = sum(aop_values.get(r, (Decimal("0"), Decimal("0")))[1] for r in refs)
            aop_values[aop] = (opening, closing)

    # Ispis redaka
    for aop, konto, opis, formula in BIL_ROWS:
        row += 1
        opening, closing = aop_values.get(aop, (Decimal("0"), Decimal("0")))

        is_section = formula and (formula.startswith("sum:") or formula.startswith("sum_konto:"))
        is_total = aop in (1, 74, 146)

        # Aktiva/Pasiva zaglavlje
        if aop == 1:
            ws.cell(row=row, column=1, value="").border = border
            ws.cell(row=row, column=2, value="A K T I V A").border = border
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
            for c in range(1, 10):
                ws.cell(row=row, column=c).font = Font(bold=True, size=11)
                ws.cell(row=row, column=c).fill = PatternFill(start_color="E8E8E8", end_color="E8E8E8", fill_type="solid")
                ws.cell(row=row, column=c).border = border
            row += 1
        elif aop == 146:
            ws.cell(row=row, column=1, value="").border = border
            ws.cell(row=row, column=2, value="P A S I V A").border = border
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
            for c in range(1, 10):
                ws.cell(row=row, column=c).font = Font(bold=True, size=11)
                ws.cell(row=row, column=c).fill = PatternFill(start_color="E8E8E8", end_color="E8E8E8", fill_type="solid")
                ws.cell(row=row, column=c).border = border
            row += 1

        ws.cell(row=row, column=1, value=konto or "").border = border
        cell_opis = ws.cell(row=row, column=2, value=opis)
        cell_opis.border = border
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)

        ws.cell(row=row, column=6, value=aop).border = border
        ws.cell(row=row, column=7, value=float(opening)).border = border
        ws.cell(row=row, column=7).number_format = "#,##0"
        ws.cell(row=row, column=8, value=float(closing)).border = border
        ws.cell(row=row, column=8).number_format = "#,##0"

        # Index
        if opening and opening != 0:
            idx = int(closing / opening * 100)
            ws.cell(row=row, column=9, value=idx).border = border
        else:
            ws.cell(row=row, column=9, value="-").border = border

        if is_total or is_section:
            for c in range(1, 10):
                ws.cell(row=row, column=c).font = bold
                if is_total:
                    ws.cell(row=row, column=c).fill = section_fill

    # Sirine stupaca
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 55
    ws.column_dimensions["F"].width = 6
    ws.column_dimensions["G"].width = 18
    ws.column_dimensions["H"].width = 18
    ws.column_dimensions["I"].width = 10

    # Save
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
